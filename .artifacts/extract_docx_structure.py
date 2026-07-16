from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document


def main() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    path = Path(sys.argv[1])
    doc = Document(path)
    payload = {
        "path": str(path),
        "paragraphs": [
            {"index": i, "style": p.style.name if p.style else None, "text": p.text}
            for i, p in enumerate(doc.paragraphs)
            if p.text.strip()
        ],
        "tables": [],
    }
    for ti, table in enumerate(doc.tables):
        payload["tables"].append(
            {
                "index": ti,
                "rows": [
                    ["\n".join(p.text for p in cell.paragraphs).strip() for cell in row.cells]
                    for row in table.rows
                ],
            }
        )
    print(json.dumps(payload, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
