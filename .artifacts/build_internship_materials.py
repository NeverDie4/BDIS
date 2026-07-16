from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(r"D:\桌面\大二学习资料\成都实训\BDIS\BDIS\8209230308-刘麒源-BDIS实训材料填写稿.docx")

NAVY = RGBColor(31, 77, 120)
BLUE = RGBColor(46, 116, 181)
MUTED = RGBColor(92, 98, 105)
BODY = RGBColor(34, 34, 34)
LIGHT = "E8EEF5"
CN_FONT = "Microsoft YaHei"
LATIN_FONT = "Calibri"


def set_run_font(run, size=None, bold=None, color=BODY, italic=None):
    run.font.name = LATIN_FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), CN_FONT)
    run._element.rPr.rFonts.set(qn("w:ascii"), LATIN_FONT)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), LATIN_FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = color


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])
    set_run_font(run, size=9, color=MUTED)


def add_body(doc, text, *, first_line=True, after=6):
    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Inches(0.33) if first_line else None
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(text)
    set_run_font(r, size=11)
    return p


def add_note(doc, label, text):
    p = doc.add_paragraph(style="Note")
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.18)
    p.paragraph_format.line_spacing = 1.2
    r = p.add_run(f"{label}：")
    set_run_font(r, size=10.5, bold=True, color=NAVY)
    r = p.add_run(text)
    set_run_font(r, size=10.5, color=MUTED)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F4F6F9")
    p_pr.append(shd)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    return p


def add_diary_page(doc, date, weekday, focus, paragraphs):
    doc.add_page_break()
    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_after = Pt(2)
    r = kicker.add_run("学生实训日志 · BDIS 项目")
    set_run_font(r, size=9.5, bold=True, color=MUTED)

    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(f"{date}（{weekday}）")
    set_run_font(r, size=16, bold=True, color=BLUE)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run(f"今日主题：{focus}")
    set_run_font(r, size=11, bold=True, color=NAVY)

    for text in paragraphs:
        add_body(doc, text)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = LATIN_FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)
    normal.font.size = Pt(11)
    normal.font.color.rgb = BODY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for level, size, before, after, color in [
        (1, 16, 18, 10, BLUE),
        (2, 13, 14, 7, BLUE),
        (3, 12, 10, 5, NAVY),
    ]:
        style = doc.styles[f"Heading {level}"]
        style.font.name = LATIN_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Note" not in [s.name for s in doc.styles]:
        note = doc.styles.add_style("Note", WD_STYLE_TYPE.PARAGRAPH)
    else:
        note = doc.styles["Note"]
    note.font.name = LATIN_FONT
    note._element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)
    note.font.size = Pt(10.5)


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    configure_styles(doc)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = hp.add_run("BDIS 实训材料填写稿 · 刘麒源")
    set_run_font(r, size=9, color=MUTED)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("第 ")
    set_run_font(r, size=9, color=MUTED)
    add_page_field(fp)
    r = fp.add_run(" 页")
    set_run_font(r, size=9, color=MUTED)

    props = doc.core_properties
    props.title = "8209230308 刘麒源 BDIS 实训材料填写稿"
    props.subject = "实习鉴定表、学生实训日志、实训体会表待填写内容"
    props.author = "刘麒源"
    props.keywords = "BDIS, 实训, 日志, 自我鉴定, 实训体会"

    # Editorial-cover opening block.
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run("软件工程专业实训材料")
    set_run_font(r, size=11, bold=True, color=MUTED)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("BDIS 项目填写稿")
    set_run_font(r, size=28, bold=True, color=NAVY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(32)
    r = p.add_run("个人总结与自我鉴定 · 7月8日至16日日志 · 实训体会")
    set_run_font(r, size=13, color=BLUE)

    for label, value in [
        ("姓名", "刘麒源"),
        ("学号", "8209230308"),
        ("学校与专业", "中南大学 · 软件工程"),
        ("项目", "BDIS 中药材数据与教学科研综合平台"),
        ("实训单位", "四川华迪信息技术有限公司"),
        ("岗位定位", "以后端工程为主，兼顾前端与全栈协作"),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(f"{label}：")
        set_run_font(r, size=10.5, bold=True, color=MUTED)
        r = p.add_run(value)
        set_run_font(r, size=10.5, color=BODY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(40)
    r = p.add_run("内容日期：2026年7月8日—2026年7月16日")
    set_run_font(r, size=9.5, italic=True, color=MUTED)

    doc.add_page_break()
    add_heading(doc, "使用说明", 1)
    add_note(
        doc,
        "复制范围",
        "本文件按三份原表的栏目分别整理。带有“填写内容”的正文可直接复制到对应表格；标题、说明和日期标签仅用于定位。",
    )
    add_note(
        doc,
        "内容依据",
        "文字以 BDIS 项目仓库、Git 提交、可访问的 Codex 对话记录和已有个人信息为依据。GPT 网页端私人聊天记录无法直接读取，因此只吸收了你在当前工作中明确体现的 GPT/Codex 协作方式，不虚构不可验证的经历。",
    )
    add_note(
        doc,
        "写作口吻",
        "整体采用学生本人第一人称，语气务实、克制，突出需求拆解、全栈协作、工程验证、复盘改进和责任意识。",
    )

    add_heading(doc, "一、实习鉴定表：个人总结与自我鉴定", 1)
    add_note(doc, "对应栏目", "“个人总结与自我鉴定”栏。以下为可直接粘贴的完整稿。")

    summary_paragraphs = [
        "在本次软件工程实训中，我主要围绕 BDIS 项目开展需求分析、前后端开发、代码审查、联调验证和版本协作。BDIS 面向中药材数据管理及教学科研业务，涉及药材资源、分布地图、生长观测、实验课程、科研项目、教学培训、评价申报、文件管理和权限控制等多个模块。面对业务范围广、角色流程多、数据关联复杂的特点，我没有只停留在完成单个页面或接口，而是先梳理项目文档、数据库模型、角色权限和模块边界，再将任务拆分为可验证的小步骤，逐步推进实现与联调。",
        "技术实践方面，我以后端工程为主线，进一步掌握了 Spring Boot、Spring Security、MyBatis-Plus、Flyway 和关系型数据库在真实项目中的协同方式。我参与梳理课程、实验记录、科研项目、培训计划及其文件访问和状态流转逻辑，关注接口参数、数据权限、迁移版本、异常处理和回归测试。同时，我也承担了较多前端与全栈协作工作，使用 Next.js、TypeScript、Ant Design 和 CSS Modules 完成首页、教学科研、地图与生长数据等界面的迭代，努力让系统既符合中药材研究场景，又保持清晰、稳定和可维护。",
        "在开发过程中，我逐渐形成了较为稳定的 AI 辅助工作方式：先使用 GPT 帮助比较方案、澄清概念和拆解复杂流程，再使用 Codex 进入实际代码库定位调用链、实施修改并运行构建和测试。我始终把 AI 当作提高检索、分析和验证效率的协作工具，而不是替代个人判断。对于登录页、导航栏、教学模块等需要保护的改动，我会先划定边界；对于评审报告中的问题，我会重新打开代码确认；对于空下拉框、构建停顿、登录失败等现象，我会优先寻找数据、环境或调用链证据，避免凭感觉大范围修改。",
        "通过多次分支同步、冲突处理、Flyway 迁移检查、格式检查、单元测试和前端构建，我更加理解了工程质量并不只体现在“功能能运行”，还体现在变更可追踪、数据可迁移、权限不越界、测试可重复以及团队成员能够安全合并。实训期间，我能够认真接受反馈，及时复盘问题，遇到陌生技术时主动查阅资料并进行小范围验证；在团队协作中，我也尽量保持沟通清楚、提交信息明确、修改范围可控。",
        "我认为自己的优势是学习速度较快、做事细致、愿意追根溯源，并且能够在前端体验、后端业务和版本管理之间建立联系。需要继续提升的地方也很明确：对大型系统的架构权衡、数据库性能调优、复杂自动化测试设计以及跨模块进度统筹还缺少足够经验。今后我会继续夯实 Java 与数据库基础，深入理解分布式系统和软件测试方法，同时保持先分析、再实现、最后验证的工作习惯，努力成长为一名既能独立解决问题、又能与团队可靠协作的软件工程师。",
    ]
    for para in summary_paragraphs:
        add_body(doc, para)

    add_heading(doc, "二、学生实训日志：2026年7月8日至7月16日", 1)
    add_note(
        doc,
        "日期说明",
        "以下包含 7 月 8 日至 16 日的全部日期。周末两天记录的是代码审查、联调准备和阶段复盘，符合项目实际连续推进的情况。",
    )
    add_body(
        doc,
        "这一阶段的日志以 BDIS 项目的真实开发轨迹为主，记录需求理解、界面实现、后端业务、权限与数据迁移、分支协作、测试验证以及使用 GPT/Codex 辅助工作的过程。",
    )

    diaries = [
        (
            "2026年7月8日",
            "星期三",
            "项目启动、业务梳理与技术基线确认",
            [
                "今天首先对 BDIS 项目的整体目标和代码结构进行了系统了解。项目包含中药材资源、分布地图、生长数据、教学科研、培训、评价申报等多个业务域，后端采用 Spring Boot、MyBatis-Plus 与 Flyway，前端采用 Next.js、TypeScript 和 Ant Design。为了避免后续开发只见局部、不见整体，我先阅读需求与设计文档，结合已有数据库迁移和实体类，整理主要角色、核心数据对象以及模块之间的关联。",
                "在 AI 辅助方面，我先用 GPT 对领域对象和常见分层方式进行对比，帮助自己理解课程、实验、科研和培训之间的区别；随后让 Codex 在实际仓库中核对目录、配置与依赖，确认结论是否符合现有代码。围绕前端，我初步确定首页不应做成普通后台管理界面，而要体现中药材研究与数字标本馆的特点，并规划了导航、主视觉搜索区、药材展示条和四个业务概览模块。",
                "今天最大的收获是认识到，项目启动阶段最重要的不是急于堆功能，而是建立一致的业务语言和技术边界。只有把需求、数据模型、权限和页面入口对应起来，后续开发才不容易反复返工。",
            ],
        ),
        (
            "2026年7月9日",
            "星期四",
            "认证权限理解与前端主框架设计",
            [
                "今天重点学习并梳理了项目的用户认证、角色权限和数据范围控制。结合后端已有的用户、角色、菜单、权限与组织机构设计，我检查了登录、Token 保存、请求拦截和当前用户信息获取的完整链路，并思考不同角色进入教学、科研、培训等模块时应看到哪些功能。通过这个过程，我对 RBAC 不再只是停留在课堂概念，而是能够把角色、权限点、接口校验和前端路由联系起来。",
                "前端方面，我继续完善系统主框架的设计思路，明确固定文本、数据文本和英文副标题的字体分工，并将首页内容划分为可独立维护的组件。对于暂时没有后端接口的数据，我采用先写稳定 mock、保留 API 边界的方式推进，避免前端直接依赖数据库结构。使用 GPT 时，我主要用它解释 Spring Security 与 JWT 的职责边界；使用 Codex 时，则让它沿请求封装、状态仓库和接口类型追踪实际调用路径。",
                "今天的工作让我体会到，权限功能看似基础，却会影响系统的每一个模块。以后设计页面和接口时，不能只考虑“能否显示”，还要同时考虑“谁可以访问、能访问哪些数据、失败后如何反馈”。",
            ],
        ),
        (
            "2026年7月10日",
            "星期五",
            "首页与教学科研界面开发、构建问题诊断",
            [
                "今天进入前端集中开发阶段，完成并整理了系统首页、药材资源和教学科研页面的阶段性成果。首页使用 HeaderNav、HeroSearchSection、HerbEssenceStrip 和 HomeOverviewGrid 组成清晰的信息层级，将药材、地图、教学和生长数据作为四个概览入口。视觉上使用较克制的纸张色、细线和中式字体，减少深阴影与过度圆角，使页面更接近中药材标本馆或研究门户，而不是通用蓝白后台。",
                "开发过程中遇到了一些细节问题：主视觉背景被多层渐变遮得过白，药材展示条最初像居中的卡片，概览区调整间距后外层高度又不一致。我逐层检查 CSS 背景和布局约束，减少遮罩，将展示条改为全宽循环滚动，并保持外层模块等高，只调整内部卡片比例。每轮改动后都通过 TypeScript 检查、lint 和构建确认没有引入新的问题。",
                "Codex 中的构建一度表现为长时间停在某个阶段，我没有把它直接判断为代码错误，而是用 GPT 梳理 Next.js 构建阶段，再对比本机 PowerShell 的调试输出，最终确认主要是执行包装和日志停顿造成的假象。今天让我更加重视“先定位问题发生在哪一层”，而不是看到异常就盲目修改业务代码。",
            ],
        ),
        (
            "2026年7月11日",
            "星期六",
            "教学科研后端设计与 JWT 认证排查",
            [
                "今天在同步最新开发分支后，重点补充了教学科研相关后端开发和数据库接口设计。围绕课程、实验记录、科研项目与培训计划，我整理了请求对象、返回对象、状态字段和数据关联，尽量让接口能够支持后续的学习、提交、批改、项目协作与培训考核流程。对于仍不确定的业务规则，我先写清假设和待确认点，没有直接把临时想法固化进代码。",
                "同时，我对登录后访问当前用户接口异常的问题进行了只读排查。通过 Codex 沿 SecurityFilterChain、JwtAuthenticationFilter、SecurityContextHolder 和 /api/auth/me 的调用链逐步定位，发现 JWT 过滤器既可能被 Servlet 容器注册，又被加入 Spring Security 过滤器链，存在重复执行风险。根据这一证据，我设计了保留安全链过滤器、禁用容器重复注册的最小修复方案，并补充相应的 SecurityConfig 回归测试思路。",
                "今天的体会是，安全问题不能靠临时绕过来解决。即使修改很小，也要理解过滤器注册机制、认证主体来源和接口上下文，并通过测试证明修复没有削弱原有的黑名单与权限校验。",
            ],
        ),
        (
            "2026年7月12日",
            "星期日",
            "阶段复盘、团队代码审查与联调准备",
            [
                "今天没有追求大范围新增功能，而是利用周末对本周代码和团队合并内容进行集中复盘。项目中评价、申报、业绩、生长记录、移动端采集和 AI 助手等模块持续加入，我按照“页面入口—后端接口—数据表—权限点—文件资源”的顺序检查它们与现有系统的连接关系，重点记录可能出现的字段不一致、状态含义冲突和数据权限遗漏。",
                "我把教学科研与培训模块的待办进一步拆分为数据库迁移、服务层规则、文件访问、前端交互和回归测试几个部分，并整理联调所需的账号、接口顺序和验证数据。GPT 主要用于帮助我比较不同工作流建模方式，Codex 则用于搜索仓库中的已有实现，避免重复造轮子或使用不一致的命名。对于不能确认的问题，我保留为检查项，准备在后续真实联调中用证据判断。",
                "这次复盘让我认识到，团队项目中“整合”本身就是重要工作。代码能单独运行并不代表放到主分支后仍然正确，提前检查边界、列出验证清单，可以显著降低后面集中修复的压力。",
            ],
        ),
        (
            "2026年7月13日",
            "星期一",
            "M12—M15 功能实现、登录与评价页面完善",
            [
                "今天根据前期设计推进 M12—M15 相关后端功能和集成材料，同时完善了登录及业务模块页面。后端工作围绕教学、实验、科研与培训的接口和数据流展开，前端则重点处理登录状态、页面入口及评价模块交互。为了保证前后端约定一致，我逐项核对请求字段、响应结构、状态值和错误处理，并补充接口使用说明，便于后续联调与团队成员复用。",
                "认证问题的最小修复也在今天得到进一步验证。我保留 JWT 过滤器在 Spring Security 链中的职责，避免修改已有认证逻辑，并通过配置与测试消除容器重复注册。随后同步远程 dev 分支，检查合并差异和潜在冲突，确保本地对登录页、导航和教学模块的改动没有被宽泛覆盖。提交前执行了差异检查和相关测试，让每次提交都能说明目的并具备基本验证证据。",
                "今天的工作让我从“完成一个功能”进一步转向“完成一个可合并、可说明、可验证的功能”。尤其是在多人协作环境中，清晰的提交记录和集成材料与代码本身同样重要。",
            ],
        ),
        (
            "2026年7月14日",
            "星期二",
            "后端评审问题修复、Flyway 与运行环境恢复",
            [
                "今天主要处理合并后的后端评审问题和运行环境故障。我先重新打开评审涉及的代码路径，检查培训流程、文件访问、权限控制和状态回退等问题是否真实存在，再进行小范围修复。随后统一处理历史 Spotless 与 Checkstyle 问题，并运行编译和测试，避免格式改动掩盖业务错误。对于来自不同分支的改动，我通过 Git 合并和差异检查确认 dev 已正确进入当前分支。",
                "数据库方面，多个 Flyway 迁移版本在合并后出现编号或校验和冲突，后端启动因此失败。我先阅读启动日志确认根因，没有把 Docker、Redis 或端口提示当成主要问题；在明确历史迁移状态无法直接兼容后，按照安全流程重建开发数据库、重新执行迁移，并验证管理账号初始化和 /api 接口可用。这个过程让我真正理解了已执行迁移的不可随意修改性，以及版本管理与数据库历史之间的关系。",
                "今天遇到的问题较多，但我始终坚持先收集证据、再决定动作。相比直接删除文件或反复重启，阅读日志、核对迁移历史和验证端口占用更有效，也让我在复杂环境故障面前更加沉着。",
            ],
        ),
        (
            "2026年7月15日",
            "星期三",
            "教学科研模块完善、工作流校准与回归测试",
            [
                "今天集中完善教学科研模块，并对教师、学生、科研和培训四类流程重新校准。教师端需要创建并发布实验课程、维护步骤与资源、批改实验报告；学生端需要选课学习、记录进度、提交报告并查看反馈；科研项目要从课程或研究方向出发，管理任务、成员、阶段成果和归档；培训则负责组合课程、项目、基地与药材并形成考核记录。通过把这些流程分别写清，我避免了将多个角色的操作混成一套简单 CRUD。",
                "前端方面继续优化教学页面的列表、操作栏和固定右侧课程详情面板，保证选中课程后按约七三比例展示，详情区域可独立滚动，关闭后恢复全宽。后端与测试方面处理了合并造成的构造参数漂移和 Flyway 版本重复，并补充了 speciesId 等关键字段的回归覆盖。一次药材下拉框为空的排查中，我先检查数据，发现启用药材数量为零，说明问题并不一定在保存逻辑。",
                "为了降低并行开发互相干扰的风险，我还整理了 Git worktree 和分支工作方式，并在合并时明确保护登录页、导航栏和教学模块。今天最大的收获是：看见界面异常时，既要检查代码，也要检查数据；看见评审意见时，既要尊重反馈，也要重新验证事实。",
            ],
        ),
        (
            "2026年7月16日",
            "星期四",
            "PR 收尾、全量验证与地图/生长模块优化",
            [
                "今天进入阶段性收尾。首先根据 PR 评审继续检查科研项目、实验记录版本、任务提交、培训流程、文件访问和 Flyway 迁移等关键路径，修复已经确认的状态流转、权限或数据关联问题。完成修改后执行格式检查、Maven 测试、Node 回归脚本和 Git 差异检查，确认修复不仅能编译，而且不会破坏已有功能。随后再次同步 dev，妥善处理冲突并保留登录、导航和教学科研模块的既有成果。",
                "在前端体验方面，我继续完善分布地图与生长数据工作台，统一两个模块的标题栏，调整页面层级、响应式断点和图表区域，使页面在不同宽度下保持清晰。为了让后续开发不再依赖口头约定，我把地图与生长模块的视觉方案、实施计划以及课程详情数据闭环整理为中文设计文档，明确哪些数据来自后端、哪些状态需要闭环、哪些页面行为必须回归验证。",
                "回顾这九天，我从项目理解、页面设计、接口实现一路做到评审修复、迁移恢复和全量验证。GPT 和 Codex 提高了我查资料、拆任务和定位代码的效率，但最终的取舍仍需要我结合业务目标、仓库现状和测试结果来判断。今天的收尾让我对自己的工作方式更加明确：先把问题说清楚，再做最小必要修改，最后用真实验证为结果负责。",
            ],
        ),
    ]
    for date, weekday, focus, paragraphs in diaries:
        add_diary_page(doc, date, weekday, focus, paragraphs)

    doc.add_page_break()
    add_heading(doc, "三、实训体会表：实训体会", 1)
    add_note(doc, "对应栏目", "实训体会正文，已超过 500 字，可直接粘贴。宣传使用是否同意请由本人在原表中勾选。")

    experience = [
        "参加本次实训之前，我对软件项目的理解更多来自课程作业：只要功能能够运行，任务就算基本完成。真正参与 BDIS 项目后，我逐渐认识到，企业级开发更重视需求是否清楚、接口是否稳定、数据是否安全、变更是否可追踪，以及团队能否在同一套规则下持续协作。四川华迪信息技术有限公司的实训安排把学习、实践、评审和复盘结合起来，让我有机会在相对真实的项目节奏中检验自己的专业能力和职业态度。",
        "BDIS 是一个围绕中药材数据及教学科研业务建设的综合平台。项目并不是简单的信息管理系统，它同时包含药材资源、分布地图、生长观测、实验课程、科研项目、培训、评价申报、文件资源和权限控制等内容。我的岗位定位以后端工程为主，但实际工作要求我不断跨越前后端边界：既要理解 Spring Boot、Spring Security、MyBatis-Plus、Flyway 和数据库迁移，也要能够阅读并修改 Next.js、TypeScript、Ant Design 与 CSS Modules 代码。正是这种完整链路的实践，让我对“全栈协作”有了更具体的认识——并不是一个人包办所有工作，而是能够理解上下游约束，减少沟通成本，让自己的改动真正融入系统。",
        "实训中给我触动最深的是工程质量意识的建立。一次登录异常，表面看是接口没有返回用户信息，实际原因却可能是 JWT 过滤器被重复注册；一次药材选择框为空，看起来像保存功能失效，最终却与启用数据为零有关；一次后端无法启动，也不能只盯着 Docker 或端口，而要从启动日志和 Flyway 校验和中确认真正原因。这些经历让我逐渐养成先观察、再假设、后验证的习惯。现在面对问题，我会优先查看日志、调用链、数据库状态和测试结果，尽量避免没有证据的大范围修改。",
        "我也在本次实训中形成了适合自己的 AI 辅助开发方法。GPT 更适合帮助我解释陌生概念、比较方案和整理业务流程，Codex 更适合进入实际代码库搜索文件、追踪调用链、执行修改并运行测试。我会先把目标、保护范围和验收标准说清楚，再让工具帮助完成重复性检索和验证。AI 的确提高了效率，但它给出的结论并不天然正确，因此我仍会重新打开代码、核对项目约定，并通过编译、单元测试、回归脚本和浏览器检查确认结果。这个过程不仅提升了我的工具使用能力，也强化了我的独立判断和信息甄别能力。",
        "团队协作方面，我体会到了版本管理的价值。项目开发期间需要频繁同步 dev 分支、处理冲突、保护本地改动、检查 Flyway 版本并准备 PR。过去我更关注代码内容本身，现在开始重视提交信息、分支状态、差异范围和验证证据。一次可靠的合并不是简单点击完成，而是要确认远程基线已经进入当前分支、关键页面没有被覆盖、数据库迁移顺序合理、测试仍然通过。通过这些实践，我的责任意识、沟通意识和风险意识都有明显提升。",
        "当然，我也看到了自己的不足。面对跨模块状态机和复杂数据权限时，我还需要更多架构设计经验；对数据库性能、自动化测试覆盖和前端响应式细节的掌握也不够深入。有时任务较多，我会在界面细节和业务闭环之间来回切换，说明自己在优先级管理上仍需加强。今后我会继续夯实 Java、数据库和计算机基础知识，主动学习软件测试、系统设计与性能优化，并坚持记录问题、复盘原因、沉淀可复用方法。",
        "总体而言，这次实训让我从“会写代码”向“能够在团队项目中负责一段完整工作”迈进了一步。我不仅提升了前后端开发能力，更学会了如何理解业务、如何与团队协作、如何借助工具提高效率，以及如何用测试和证据对结果负责。这些收获将成为我今后学习和求职的重要基础。",
    ]
    for para in experience:
        add_body(doc, para)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(18)
    r = p.add_run("刘麒源\n2026年7月16日")
    set_run_font(r, size=11, bold=True, color=NAVY)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
